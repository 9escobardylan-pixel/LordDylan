#!/usr/bin/env python3
"""Generates Dylan Escobar's ATS-friendly resume as a .docx file."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0xE8, 0x40, 0x2C)
INK = RGBColor(0x0E, 0x0E, 0x10)
FONT = "Calibri"

doc = Document()

# ---- Page setup: single column, tight-but-readable margins ----
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

# ---- Base style ----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.05
rPr = normal.element.get_or_add_rPr()
rFonts = rPr.find(qn("w:rFonts"))
if rFonts is None:
    rFonts = OxmlElement("w:rFonts")
    rPr.append(rFonts)
rFonts.set(qn("w:eastAsia"), FONT)


def set_cell_margins():
    pass


def add_bottom_border(paragraph, color="E8402C", size=6, space=2):
    """Adds a thin bottom rule to a paragraph (used under section headers)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_run(paragraph, text, bold=False, italic=False, color=INK, size=10.5, caps=False):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.all_caps = caps
    return run


def section_header(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    add_run(p, text, bold=True, color=ACCENT, size=11.5, caps=True)
    # widen letter spacing slightly for a header feel via a trailing space trick is unreliable;
    # rely on caps + bold + accent + rule instead.
    add_bottom_border(p)
    return p


def body_paragraph(text, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    add_run(p, text)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.2)
    pf = p.paragraph_format
    pf.line_spacing = 1.05
    add_run(p, text)
    return p


def job_header(title_org, dates):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    tab_stops = p.paragraph_format.tab_stops
    usable_width = section.page_width - section.left_margin - section.right_margin
    tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
    add_run(p, title_org, bold=True, size=10.5)
    add_run(p, "\t")
    add_run(p, dates, italic=True, size=10)
    return p


# ================= HEADER: NAME + CONTACT =================
name_p = doc.add_paragraph()
name_p.paragraph_format.space_after = Pt(2)
add_run(name_p, "DYLAN ESCOBAR", bold=True, color=ACCENT, size=26)

title_p = doc.add_paragraph()
title_p.paragraph_format.space_after = Pt(4)
add_run(title_p, "Freelance Direct-Response Creative Strategist", size=11.5, color=INK, italic=True)

contact_p = doc.add_paragraph()
contact_p.paragraph_format.space_after = Pt(0)
add_run(contact_p, "dylan@dylanescobar.com", size=10)
add_run(contact_p, "   |   ", size=10, color=ACCENT)
add_run(contact_p, "1-661-525-8048", size=10)

# thin divider under the whole header block
add_bottom_border(contact_p, size=8, space=6)

# ================= SUMMARY =================
section_header("Summary")
body_paragraph(
    "Freelance creative strategist specializing in Meta ad hook/script writing, static and "
    "AI-video ad production, and ad research for DTC brands. Built and grew a YouTube channel "
    "(@BillionaireHustler) to 700K+ lifetime views and 800+ subscribers with zero ad spend "
    "(91%+ like ratio, 42K-view top clip) — applying direct-response and audience-retention "
    "principles to client ad work."
)

# ================= EXPERIENCE =================
section_header("Experience")
job_header("Freelance Creative Strategist — Self-Employed (Remote)", "May 2026 – June 2026")
bullet(
    "Engaged by a multi-brand DTC portfolio (Tailored Canvases, WonderMe, TellMyTale) to build "
    "full ad script systems across three brands in parallel."
)
bullet(
    "Tailored Canvases: Delivered three production-ready ad scripts with beat-by-beat structure "
    "and optimization notes for a home decor DTC brand."
)
bullet(
    "WonderMe: Wrote three production-ready scripts including 36 hook variations for Meta ad testing."
)
bullet(
    "TellMyTale: Built full ad script systems across two product lines (Dinosaurs, Superhero); "
    "managed a live $330 test spend achieving a 40.80% new-visitor rate and 1.5 blended ROAS."
)
bullet(
    "Resolved a client payment dispute independently via Upwork's milestone system, securing "
    "full payment without conceding to an off-platform reduced settlement."
)

# ================= ADDITIONAL PROJECTS =================
section_header("Additional Projects")
bullet("Brand/VOC research for Enhanced Human, Outdoor Vitals, Beam, ManTalks, and Hatch.")
bullet(
    "Ad creative and copy across DTC verticals: electrolyte drinks, travel gear, cosmetics, "
    "wellness, and consumer products (spec + client work)."
)

# ================= EDUCATION / TRAINING =================
section_header("Education / Training")
bullet("Copy Millions Blueprint — direct-response copywriting and creative strategy coaching program (current)")
bullet("Eagle Scout")

# ================= SKILLS =================
section_header("Skills")
body_paragraph(
    "Meta ad strategy, direct-response copywriting, hook writing, script writing, AI video ad "
    "production, static ad design, VOC/audience research, YouTube content strategy, DTC brand research",
    space_after=0,
)

doc.save("/home/user/LordDylan/resume/Dylan_Escobar_Resume.docx")
print("Saved resume.")
